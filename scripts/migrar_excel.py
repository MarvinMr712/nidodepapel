# -*- coding: utf-8 -*-
"""Migra el Excel de inventario a un dataset tipado con la taxonomia limpia."""
import json, re, unicodedata
from collections import defaultdict
import openpyxl

SRC = r"C:\Users\User\Downloads\Inventario_Cuentos_Infantiles 2026.xlsx"
OUT = r"C:\repositorios\casa-de-cuentos\src\data\libros.json"

# ---------------------------------------------------------------- utilidades
def strip_accents(s):
    return "".join(c for c in unicodedata.normalize("NFD", str(s))
                   if unicodedata.category(c) != "Mn")

def norm_key(s):
    return re.sub(r"[^a-z0-9]", "", strip_accents(str(s)).lower())

def slugify(s):
    s = strip_accents(str(s)).lower()
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return re.sub(r"^-+|-+$", "", s)

def clean_text(s):
    s = unicodedata.normalize("NFC", str(s))
    s = s.replace("\u00a0", " ").replace("\ufffd", "")
    s = re.sub(r"\s+", " ", s).strip()
    s = s.lstrip("?!.,;:").strip()          #guildas de interrogacion/exclamacion sueltas
    return s

# ---------------------------------------------------------------- taxonomias
CATEGORIAS = {
    "emociones-vinculo":       "Emociones y vínculo",
    "habitos-rutinas-valores": "Hábitos, rutinas, valores y límites",
    "conocimiento":            "Conocimiento",
    "diversion-imaginacion":   "Diversión, imaginación, creatividad",
    "fe":                      "Fe",
}
CAT_POR_FILTRO = {
    "emociones y vinculo": "emociones-vinculo",
    "habitos, rutinas, valores y limites": "habitos-rutinas-valores",
    "conocimiento": "conocimiento",
    # los libros en ingles se cuelgan en Conocimiento (animales, vehiculos, ciencias, espacio, ingles)
    "diversion, imaginacion, creatividad (ingles)": "conocimiento",
    "diversion, imaginacion, creatividad": "diversion-imaginacion",
    "fe": "fe",
}

AUTOR_CANON = {
    "disney": "Disney",
    "lexus editorial": "Editorial Lexus",
    "lexus editores": "Editorial Lexus",
    "editorial lexus": "Editorial Lexus",
    "varios / editorial": "Varios",
    "varios": "Varios",
    "susaeta / varios": "Susaeta y Varios",
    "pichoncito": "Pichoncito",
    "libsa": "Libsa",
    "fondo de cultura economico": "Fondo de Cultura Economica",
}

def canon_autor(raw):
    a = clean_text(raw)
    return AUTOR_CANON.get(a.lower(), a)

FORMATO_MAT = [
    ("carton grueso", "Cartón grueso"),
    ("carton acolchado", "Cartón grueso"),
    ("desplegable", "Pop-up"),
    ("pop up", "Pop-up"),
    ("pop-up", "Pop-up"),
    ("carton", "Cartón"),
    ("delgada", "Tapa delgada"),
    ("tapa dura", "Tapa dura"),
]

def parse_formato(raw):
    t = strip_accents(clean_text(raw)).lower()
    if not t:
        return "Tapa dura", False, ""
    material = "Tapa dura"
    for key, val in FORMATO_MAT:
        if key in t:
            material = val
            break
    inter = any(k in t for k in ("interactiv", "mecanismo", "solapa", "pop",
                                 "despleg", "puppet", "musical", "musica", "juguete", "peque"))
    return material, inter, clean_text(raw)

def parse_edad(raw):
    """'4 - 8 anos' -> (4, 8). Devuelve (None, None) si no se puede."""
    t = clean_text(raw).lower()
    m = re.findall(r"(\d+)", t)
    if len(m) >= 2:
        lo, hi = int(m[0]), int(m[1])
        if lo <= hi:
            return lo, hi
    if len(m) == 1:
        v = int(m[0])
        return v, v + 2
    return None, None

ETAPAS = [
    {"slug": "primeros-anos", "nombre": "Primeros años",
     "min": 0, "max": 3,
     "descripcion": "El vínculo marca la diferencia: libros para leer juntos, con voz alta y en casa."},
    {"slug": "pequenos-lectores", "nombre": "Pequeños lectores",
     "min": 4, "max": 6,
     "descripcion": "Empiezan a elegir, a repetir sus favoritos y a preguntar por qué."},
    {"slug": "lectores-en-evolucion", "nombre": "Lectores en evolución",
     "min": 7, "max": 9,
     "descripcion": "Sostienen historias más largas, buscan temas propios y construyen hábito."},
]

def etapas_de(edad_min, edad_max):
    """Un libro sirve a toda etapa con la que su rango de edad se solapa.
    '3 - 5 anos' sirve a Primeros anos (3) y a Pequenos lectores (5)."""
    if edad_min is None and edad_max is None:
        return [ETAPAS[0]["slug"]]
    lo = edad_min if edad_min is not None else edad_max
    hi = edad_max if edad_max is not None else edad_min
    hits = []
    for e in ETAPAS:
        e_hi = e["max"] if e["max"] is not None else 99
        if lo <= e_hi and hi >= e["min"]:
            hits.append(e["slug"])
    return hits or [ETAPAS[0]["slug"]]

# ---------------------------------------------------------------- lectura
wb = openpyxl.load_workbook(SRC, data_only=True)
sheet = [n for n in wb.sheetnames if "Cat" in n][0]
rows = list(wb[sheet].iter_rows(values_only=True))
raw = [r for r in rows[4:] if any(c is not None and str(c).strip() for c in r)]

grupos = defaultdict(list)
for r in raw:
    grupos[norm_key(r[1])].append(r)

libros, copias_total, duplicados = [], 0, 0
for key, filas in grupos.items():
    if len(filas) > 1:
        duplicados += 1
    base = filas[0]
    titulo = clean_text(base[1])
    material, inter, fmt_raw = parse_formato(base[5])
    e_min, e_max = parse_edad(base[3])
    cat_filtro = strip_accents(clean_text(base[7])).lower()
    cat = CAT_POR_FILTRO.get(cat_filtro, "diversion-imaginacion")
    autor = canon_autor(base[2])

    copias = []
    for i, f in enumerate(filas, start=1):
        copias.append({
            "codigo": f"{str(f[0]).zfill(3)}",
            "propietario": clean_text(f[6]),
            "estado": "disponible",
        })
    copias_total += len(copias)

    libros.append({
        "id": f"lib_{str(base[0]).zfill(3)}",
        "slug": slugify(titulo),
        "titulo": titulo,
        "autor": autor,
        "tema": clean_text(base[4]),          # texto libre: descriptivo, NO filtro
        "categoria": cat,
        "formato": material,
        "interactivo": inter,
        "formatoOriginal": fmt_raw,
        "edadMin": e_min,
        "edadMax": e_max,
        "etapas": etapas_de(e_min, e_max),
        "copias": copias,
    })

# slugs unicos
vistos = defaultdict(int)
for l in libros:
    base = l["slug"] or "libro"
    vistos[base] += 1
    l["slug"] = base if vistos[base] == 1 else f"{base}-{vistos[base]}"

libros.sort(key=lambda l: l["titulo"].lower())

payload = {
    "categorias": [{"slug": k, "nombre": v} for k, v in CATEGORIAS.items()],
    "etapas": ETAPAS,
    "libros": libros,
}

with open(OUT, "w", encoding="utf-8") as f:
    json.dump(payload, f, ensure_ascii=False, indent=2)

# ---------------------------------------------------------------- reporte
print(f"titulos unicos : {len(libros)}")
print(f"copias fisicas : {copias_total}")
print(f"grupos con >1 copia: {duplicados}")
print("etapas (un libro puede estar en varias):")
for e in ETAPAS:
    print(f"   {e['nombre']:<24} {sum(1 for l in libros if e['slug'] in l['etapas'])}")
print("solo en una etapa:", sum(1 for l in libros if len(l["etapas"]) == 1))
print("categorias     : " + ", ".join(
    f"{v}={sum(1 for l in libros if l['categoria']==k)}" for k, v in CATEGORIAS.items()))
sin_edad = [l["titulo"] for l in libros if l["edadMin"] is None]
print(f"sin edad parseable: {len(sin_edad)}")
for t in sin_edad[:10]:
    print("   ", t)
print(f"formatos       : " + ", ".join(
    f"{k}={sum(1 for l in libros if l['formato']==k)}"
    for k in sorted({l['formato'] for l in libros})))
print(f"interactivos   : {sum(1 for l in libros if l['interactivo'])}")
